import time

from app import appraisal, dedupe, model

SCHEME = appraisal.normalize_scheme({'indicators': [
    {'id': 'I1', 'group': '党建', 'name': '学生党建工作', 'max': 10, 'method': 'judge',
     'criteria': [{'level': '优秀', 'score': 10, 'description': '计划+3次活动'}, {'level': '良好', 'score': 8, 'description': '计划+1-2次'}]},
    {'id': 'I3', 'group': '日常', 'name': '工作记录', 'max': 15, 'method': 'checklist', 'required': ['辅导员工作日志', '谈话记录', '值班记录'], 'deduct': 5},
    {'id': 'I5', 'group': '拔尖', 'name': '学生竞赛获奖', 'max': 10, 'method': 'count', 'points': {'省级及以上': 2}, 'cap': 10},
    {'id': 'I7', 'group': '心理', 'name': '心理工作队伍', 'max': 5, 'method': 'checklist', 'required': ['年级心理委员名单', '心理健康普测组织材料'], 'deduct': 3},
    {'id': 'I9', 'group': '队伍', 'name': '理论研究成果', 'max': 15, 'method': 'count', 'points': {'核心期刊': 5, 'CN期刊': 3}, 'cap': 15},
]})['indicators']


def material(mid, person, indicator, **extra):
    return {'id': mid, 'person': person, 'indicator': indicator, 'file': mid + '.txt', 'summary': '', 'quotes': [], **extra}


def test_count_merges_same_project_and_respects_levels():
    materials = [
        material('M1', '王敏', 'I5', event='挑战杯', project='智慧校园能耗监测系统', level='省级'),
        material('M2', '王敏', 'I5', event='挑战杯', project='《智慧校园能耗监测系统》', level='省级二等奖'),
        material('M3', '王敏', 'I5', event='数学建模', project='共享单车调度', level='全国一等奖'),
        material('M4', '王敏', 'I5', event='校程序设计赛', project='', level='校级'),
        material('M5', '王敏', 'I9', project='谈心谈话的实践路径', level='中文核心期刊'),
        material('M6', '王敏', 'I9', project='谈心谈话的实践路径', level='CN期刊'),
    ]
    items = {i['id']: i for i in appraisal.score_person('王敏', SCHEME, materials)['items']}
    assert items['I5']['score'] == 4  # two certificates of one project count once; 校级 is out of range
    assert any('2 份材料合并为 1 项' in items['I5']['how'] for _ in [0]) and '级别不在计分范围' in items['I5']['how']
    assert items['I9']['score'] == 5  # same paper in two journals: highest only


def test_checklist_deducts_and_never_goes_negative():
    materials = [material('M1', '李强', 'I3', doc_type='辅导员工作日志'), material('M2', '李强', 'I3', doc_type='值班记录')]
    items = {i['id']: i for i in appraisal.score_person('李强', SCHEME, materials)['items']}
    assert items['I3']['score'] == 10 and items['I3']['missing'] == ['谈话记录']
    assert items['I7']['score'] == 0  # both missing: 5 - 6 floors at 0


def test_judge_waits_for_evaluator_and_agreement_ignores_empty_suggestions():
    result = appraisal.score_person('王敏', SCHEME, [], suggestions={'王敏::I1': {'level': '', 'score': None}})
    assert result['pending'] == ['I1'] and result['total'] is None
    judged = appraisal.score_person('王敏', SCHEME, [], {'王敏::I1': {'score': 10}}, {'王敏::I1': {'level': '良好', 'score': 8}})
    other = appraisal.score_person('李强', SCHEME, [], {'李强::I1': {'score': 0}}, {'李强::I1': {'level': '', 'score': None}})
    assert judged['total'] == 10 and appraisal.agreement([judged, other]) == {'suggested': 1, 'kept': 0}


def test_flags_cross_claims_only_for_counted_items():
    materials = [
        material('M1', '王敏', 'I5', project='能耗监测', level='省级'), material('M2', '李强', 'I5', project='能耗监测', level='省级'),
        material('M3', '王敏', 'I7', doc_type='年级心理委员名单', event='心理委员名单'), material('M4', '赵琳', 'I7', doc_type='年级心理委员名单', event='心理委员名单'),
        material('M5', '', 'I3'), material('M6', '王敏', 'I5', project='程序设计', level='校级'),
    ]
    flags = {m['id']: m['flags'] for m in appraisal.material_flags(materials, ['王敏', '李强', '赵琳'], SCHEME)}
    assert any('他人申报' in f for f in flags['M1']) and any('他人申报' in f for f in flags['M2'])
    assert flags['M3'] == [] and flags['M4'] == []
    assert any('属于哪位' in f for f in flags['M5'])
    assert any('计分范围' in f for f in flags['M6'])


def test_checklist_material_with_unmatched_type_is_flagged():
    materials = [material('M1', '赵琳', 'I7', doc_type='心理委员名册'), material('M2', '赵琳', 'I7', doc_type='年级心理委员名单')]
    flags = {m['id']: m['flags'] for m in appraisal.material_flags(materials, ['赵琳'], SCHEME)}
    assert any('对不上必交材料' in f for f in flags['M1']) and flags['M2'] == []


def test_flat_points_for_a_level_rule_are_warned():
    scheme = appraisal.normalize_scheme({'indicators': [{'name': '学生竞赛获奖', 'max': 10, 'method': 'count', 'points': 2, 'cap': 10}],
                                         'source_text': '5. 学生竞赛获奖（10分）：所带学生获省级及以上竞赛奖励，每个项目计2分，最高10分。'})
    assert any('没有按级别填写' in w for w in appraisal.check_scheme(scheme)[1])


def test_quote_check_marks_invented_quotes():
    quotes = appraisal.verify_quotes(['全年开展谈心谈话 214 人次', '获得全国优秀辅导员称号'], ['全年开展谈心谈话 214 人次，其中……'])
    assert [q['verified'] for q in quotes] == [True, False]


def fake_model(instruction, payload, images=None):
    """Test stand-in for the model adapter; it reads the synthetic file names, not the content."""
    if instruction == appraisal.SCHEME_INSTRUCTION:
        return {'indicators': SCHEME}
    if instruction == appraisal.REVISE_INSTRUCTION:  # “科技讲座每次 3 分” style edit: here the 竞赛 points change
        return {'indicators': [{**i, 'points': {'省级及以上': 3}} if i['id'] == 'I5' else i for i in payload['current']['indicators']]}
    if instruction == appraisal.JUDGE_INSTRUCTION:
        return {'level': '良好', 'reason': '测试', 'quotes': []}
    if instruction == dedupe.GROUP_INSTRUCTION:  # one achievement per title; within it one award per level
        works = {}
        for c in payload['cards']:
            works.setdefault(c['project'], {}).setdefault(c['level'], []).append(c['id'])
        return {'achievements': [{'awards': list(levels.values()), 'reason': '测试', 'evidence': []}
                                 for levels in works.values() if sum(len(x) for x in levels.values()) > 1]}
    name = payload['file_name']
    person = payload['filename_hint']
    for words, indicator, doc_type in [(('日志',), 'I3', '辅导员工作日志'), (('谈话',), 'I3', '谈话记录'), (('值班',), 'I3', '值班记录'),
                                       (('证书',), 'I5', '获奖证书'), (('论文',), 'I9', '论文'), (('心理委员',), 'I7', '年级心理委员名单'),
                                       (('普测',), 'I7', '心理健康普测组织材料'), (('党',), 'I1', '党建记录')]:
        if any(w in name for w in words):
            level = '校级' if '校程序' in name else '核心期刊' if '核心' in name else 'CN期刊' if 'CN' in name else '国家级' if '国赛' in name else '省级'
            project = '能耗监测' if '挑战杯' in name else '谈心谈话' if '论文' in name else name
            return {'person': person, 'doc_type': doc_type, 'indicators': [{'id': indicator, 'reason': '测试'}], 'level': level, 'project': project, 'summary': name}
    return {'person': person, 'indicators': [], 'summary': name}


def wait(client, task_id):
    for _ in range(300):
        task = client.get('/api/tasks/' + task_id).json()
        if task['status'] not in ('queued', 'running'):
            return task
        time.sleep(.03)
    raise AssertionError('Task did not settle')


def test_appraisal_demo_end_to_end(monkeypatch):
    from fastapi.testclient import TestClient
    from app.main import app
    monkeypatch.setattr(model, 'ask_json', fake_model)
    monkeypatch.setattr(model, 'status', lambda: {'configured': True, 'base_url': 'test', 'model': 'test', 'key_configured': True})
    with TestClient(app) as client:
        tid = client.post('/api/demo/appraisal', json={}).json()['id']
        task = wait(client, tid)
        pending = task['result']['pending']
        assert task['kind'] == 'appraisal' and pending['stage'] == 'confirm' and pending['detected']['kind'] == 'appraisal'
        assert set(pending['people']) == {'王敏', '李强', '赵琳'}  # no list was typed in: taken from the file names
        assert len(pending['materials']) == 30 and len(pending['scheme_lines']) == 5 and len(pending['judge']) == 3
        # An edit to the indicators sends the materials back to be read against the new scheme.
        revised = client.post(f'/api/tasks/{tid}/revise', json={'text': '竞赛获奖每项改为3分'}).json()
        assert revised['scheme_changed'] and '3 分' in revised['scheme_lines'][2]
        assert client.post(f'/api/tasks/{tid}/review', json={}).status_code == 200
        task = wait(client, tid)
        pending = task['result']['pending']
        assert pending['stage'] == 'confirm' and not pending.get('scheme_changed')
        changes = {m['id']: {'person': m['person'], 'indicator': m['indicator'], 'excluded': m['file'] == '李强_获奖证书_挑战杯省赛.png'} for m in pending['materials']}
        missing = client.post(f'/api/tasks/{tid}/review', json={'materials': changes, 'scheme_answers': {'I9': 'highest'}})
        assert missing.status_code == 400 and '需要你打分' in missing.json()['detail']
        assert client.post(f'/api/tasks/{tid}/review', json={'materials': changes, 'judgements': {'王敏::I1': {'score': 11}}, 'scheme_answers': {'I9': 'highest'}}).status_code == 400
        judgements = {f'{p}::I1': {'score': 8} for p in ('王敏', '李强', '赵琳')}
        assert client.post(f'/api/tasks/{tid}/review', json={'materials': changes, 'judgements': judgements, 'scheme_answers': {'I9': 'highest'}}).status_code == 200
        task = wait(client, tid)
        assert task['status'] == 'completed', task
        totals = {r['person']: r for r in task['result']['results']}
        items = {i['id']: i['score'] for i in totals['王敏']['items']}
        assert items['I5'] == 6 and items['I3'] == 15 and items['I7'] == 2  # two projects at the revised 3 points
        assert {i['id']: i['score'] for i in totals['李强']['items']}['I9'] == 5
        assert client.get(f'/api/tasks/{tid}/download/summary.csv').status_code == 200


def test_appraisal_new_version_keeps_scores_unless_their_materials_changed(monkeypatch):
    import io
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from app.main import app
    asked = []
    monkeypatch.setattr(model, 'ask_json', lambda instruction, payload, images=None: asked.append(instruction) or fake_model(instruction, payload, images))
    monkeypatch.setattr(model, 'status', lambda: {'configured': True, 'base_url': 'test', 'model': 'test', 'key_configured': True})
    with TestClient(app) as client:
        tid = client.post('/api/demo/appraisal', json={}).json()['id']
        pending = wait(client, tid)['result']['pending']
        changes = {m['id']: {'excluded': True, 'note': '重复'} for m in pending['materials'] if m['file'] == '李强_获奖证书_挑战杯省赛.png'}
        judgements = {f'{p}::I1': {'score': 8} for p in ('王敏', '李强', '赵琳')}
        assert client.post(f'/api/tasks/{tid}/review', json={'materials': changes, 'judgements': judgements, 'scheme_answers': {'I9': 'highest'}}).status_code == 200
        assert wait(client, tid)['status'] == 'completed'
        before = asked.count(appraisal.SCHEME_INSTRUCTION)
        added = ('files', ('王敏_党建活动记录_补充.txt', io.BytesIO('王敏 组织党日活动'.encode('utf-8')), 'text/plain'))
        assert client.post(f'/api/tasks/{tid}/modify', files=[added]).status_code == 202
        task = wait(client, tid)
        pending = task['result']['pending']
        assert asked.count(appraisal.SCHEME_INSTRUCTION) == before  # the confirmed indicators are carried
        assert [x['id'] for x in pending['changes']['lapsed']] == ['王敏::I1']  # scored on other materials than now
        assert {j['person'] for j in pending['judge'] if not j['decided']} == {'王敏'}
        assert next(m for m in pending['materials'] if m['file'] == '李强_获奖证书_挑战杯省赛.png')['excluded']  # the correction is kept
        book = load_workbook(io.BytesIO(client.get(f'/api/tasks/{tid}/download/results.xlsx').content))
        assert book.sheetnames[:2] == ['说明（暂定）', '汇总（暂定）'] and '待评分' in book.sheetnames
