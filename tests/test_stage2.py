"""Changing a task at any stage: new versions, decisions kept or lapsed, undo, stop; result states, the workbook, locating a record."""
import io
import json
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import model, scoring, store, understand, versions, workflow
from app.main import app

CONFIGURED = {'configured': True, 'base_url': 'test', 'model': 'test', 'key_configured': True}
ROSTER = '学号,姓名\n001,甲\n002,乙\n003,丙\n'
FINAL = '学号,姓名,考核项目,成绩\n001,甲,期末,80\n002,乙,期末,90\n003,丙,期末,70\n'
USUAL = '学号,姓名,考核项目,成绩\n001,甲,平时,85\n002,乙,平时,75\n'
EXTRA = '学号,姓名,考核项目,成绩\n001,甲,期末,82\n'
RULE = {'components': [{'name': '期末', 'assessments': ['期末'], 'weight': 60, 'missing': 'zero'},
                       {'name': '平时', 'assessments': ['平时'], 'weight': 40, 'missing': 'zero'}], 'rounding': {'digits': 1}}


def csv(name, body):
    return ('files', (name, io.BytesIO(body.encode('utf-8')), 'text/csv'))


def mocked(monkeypatch):
    calls = {'parse': 0}

    def ask(instruction, payload, images=None, fresh=False):
        if instruction == understand.PLAN_INSTRUCTION:
            return {'subject': '学生', 'questions': []}
        if instruction == scoring.PARSE_INSTRUCTION:
            calls['parse'] += 1
            return RULE
        if instruction == scoring.REVISE_INSTRUCTION:
            return {**RULE, 'rounding': {'digits': 0}, 'explanation': '改为整数'}
        return {}
    monkeypatch.setattr(model, 'ask_json', ask)
    monkeypatch.setattr(model, 'status', lambda: CONFIGURED)
    return calls


def wait(client, tid, settled=('waiting_review', 'completed', 'failed', 'stopped')):
    for _ in range(400):
        task = client.get('/api/tasks/' + tid).json()
        if task['status'] in settled:
            return task
        time.sleep(.03)
    raise AssertionError('Task did not settle')


def create(client, *files):
    response = client.post('/api/tasks', files=[csv('名单.csv', ROSTER), *files], data={'kind': 'grades', 'rules_text': '期末60%，平时40%，没有记录按0分。'})
    assert response.status_code == 201, response.text
    return response.json()['id']


def decide(client, tid, task, pick=lambda options: options[-1]):
    """Answer every open problem: a conflict takes the chosen record, anything else its suggestion."""
    decisions = {}
    for i in task['result']['pending']['issues']:
        decisions[i['id']] = {'action': 'select', 'record_id': pick(i['options'])['value']} if i['type'] == 'conflict' else i['suggestion']
    response = client.post(f'/api/tasks/{tid}/review', json={'decisions': decisions})
    assert response.status_code == 200, response.text
    return wait(client, tid)


def modify(client, tid, files=(), **data):
    response = client.post(f'/api/tasks/{tid}/modify', files=list(files) or None,
                           data={k: json.dumps(v) if isinstance(v, list) else v for k, v in data.items()})
    assert response.status_code == 202, response.text
    return wait(client, tid)


def totals(task):
    return {t['student_id']: t['total'] for t in task['result']['totals']}


def test_adding_material_keeps_decisions_and_moves_only_affected_totals(monkeypatch):
    calls = mocked(monkeypatch)
    with TestClient(app) as client:
        tid = create(client, csv('期末.csv', FINAL), csv('平时.csv', USUAL), csv('补充.csv', EXTRA))
        task = wait(client, tid)
        assert task['state']['code'] == 'needs_review' and [i['type'] for i in task['result']['pending']['issues']] == ['conflict']
        task = decide(client, tid, task, pick=lambda o: next(x for x in o if '82' in x['label']))
        assert task['status'] == 'completed' and task['state']['code'] == 'confirmed'
        assert totals(task) == {'001': 83.2, '002': 84.0, '003': 42.0}

        task = modify(client, tid, [csv('平时补交.csv', '学号,姓名,考核项目,成绩\n003,丙,平时,60\n')])
        pending = task['result']['pending']
        assert task['version'] == 2 and task['status'] == 'waiting_review' and not pending['issues']  # the conflict is not asked again
        changes = pending['changes']
        assert changes['was_confirmed'] and changes['kept'] == 1 and not changes['lapsed'] and changes['files_added'] == ['平时补交.csv']
        assert [(t['student_id'], t['before'][0], t['after'][0]) for t in changes['totals_moved']] == [('003', 42.0, 66.0)]
        assert changes['totals_same'] == 2 and calls['parse'] == 1  # the confirmed rule is carried, not read again
        assert [d['label'] for d in pending['decided']] and '82' in pending['decided'][0]['label']
        task = decide(client, tid, task)
        assert totals(task) == {'001': 83.2, '002': 84.0, '003': 66.0}

        # The same materials and the same decision from scratch give the same result.
        fresh = create(client, csv('期末.csv', FINAL), csv('平时.csv', USUAL), csv('补充.csv', EXTRA),
                       csv('平时补交.csv', '学号,姓名,考核项目,成绩\n003,丙,平时,60\n'))
        again = decide(client, fresh, wait(client, fresh), pick=lambda o: next(x for x in o if '82' in x['label']))
        assert totals(again) == totals(task)


def test_a_decision_lapses_when_its_problem_changes_and_undo_restores(monkeypatch):
    mocked(monkeypatch)
    with TestClient(app) as client:
        tid = create(client, csv('期末.csv', FINAL), csv('平时.csv', USUAL), csv('补充.csv', EXTRA))
        task = decide(client, tid, wait(client, tid), pick=lambda o: next(x for x in o if '82' in x['label']))
        task = modify(client, tid, [csv('期末更正.csv', '学号,姓名,考核项目,成绩\n001,甲,期末,85\n')])
        pending = task['result']['pending']
        # A third score for the same exam: the earlier choice between two no longer answers it.
        assert len(pending['changes']['lapsed']) == 1 and pending['issues'][0]['was'] and len(pending['issues'][0]['options']) == 3
        assert task['state']['code'] == 'needs_review'
        assert client.get(f'/api/tasks/{tid}/download/totals.csv').status_code == 404  # the confirmed result is not served as current
        assert client.post(f'/api/tasks/{tid}/undo').json()['version'] == 1
        task = client.get('/api/tasks/' + tid).json()
        assert task['status'] == 'completed' and totals(task)['001'] == 83.2 and task['version'] == 1
        assert client.get(f'/api/tasks/{tid}/download/totals.csv').status_code == 200
        assert client.post(f'/api/tasks/{tid}/undo').status_code == 400  # nothing before the first version


def test_removing_a_file_makes_its_problem_moot_and_revoking_asks_again(monkeypatch):
    mocked(monkeypatch)
    with TestClient(app) as client:
        tid = create(client, csv('期末.csv', FINAL), csv('平时.csv', USUAL), csv('补充.csv', EXTRA))
        task = decide(client, tid, wait(client, tid))
        conflict = task['result']['decided'][0]['id']
        task = modify(client, tid, revoke=[conflict])
        pending = task['result']['pending']
        assert [i['id'] for i in pending['issues']] == [conflict] and not pending['issues'][0]['was']
        assert pending['changes']['revoked'][0]['id'] == conflict
        index = next(f['index'] for f in task['files'] if f['name'] == '补充.csv')
        task = modify(client, tid, remove=[index])
        pending = task['result']['pending']
        assert not pending['issues'] and pending['changes']['files_removed'] == ['补充.csv']
        assert {t['student_id']: t['total'] for t in pending['preview']}['001'] == 82.0  # 80×60% + 85×40%
        assert client.post(f'/api/tasks/{tid}/modify', data={}).status_code == 400  # nothing changed


def test_identical_upload_is_not_a_change_and_the_same_name_replaces(monkeypatch):
    mocked(monkeypatch)
    with TestClient(app) as client:
        tid = create(client, csv('期末.csv', FINAL), csv('平时.csv', USUAL))
        task = decide(client, tid, wait(client, tid))
        response = client.post(f'/api/tasks/{tid}/modify', files=[csv('平时副本.csv', USUAL)])
        assert response.status_code == 400 and '内容相同' in response.json()['detail']
        task = modify(client, tid, [csv('平时.csv', USUAL.replace('75', '95'))])
        assert task['result']['pending']['changes']['files_replaced'] == ['平时.csv']
        assert sum(1 for f in task['files'] if f['name'] == '平时.csv') == 1
        task = decide(client, tid, task)
        assert totals(task)['002'] == 92.0


def test_rule_edits_survive_a_new_version(monkeypatch):
    calls = mocked(monkeypatch)
    with TestClient(app) as client:
        tid = create(client, csv('期末.csv', FINAL), csv('平时.csv', USUAL))
        wait(client, tid)
        assert client.post(f'/api/tasks/{tid}/revise', json={'text': '总评取整数'}).status_code == 200
        task = modify(client, tid, [csv('平时补交.csv', '学号,姓名,考核项目,成绩\n003,丙,平时,61\n')])
        pending = task['result']['pending']
        assert pending['rule']['rounding']['digits'] == 0 and pending['revisions'][0]['text'] == '总评取整数' and calls['parse'] == 1
        task = modify(client, tid, rules_text='期末50%，平时50%，没有记录按0分。')
        assert task['result']['pending']['changes']['rules_changed'] and calls['parse'] == 2  # changed rules are read again


def test_stop_between_steps_and_continue(monkeypatch, tmp_path):
    monkeypatch.setattr(model, 'status', lambda: {**CONFIGURED, 'configured': False})
    (tmp_path / '期末.csv').write_text(FINAL, encoding='utf-8')
    tid = store.create_task('中止测试', 'grades', {'rules_text': '', 'course': '', 'use_model': False}, [{'name': '期末.csv', 'path': str(tmp_path / '期末.csv'), 'role': ''}])
    versions.request_stop(tid)
    workflow.run(tid)
    task = store.get_task(tid)
    assert task['status'] == 'stopped' and not versions.stopping(tid)
    workflow.run(tid)  # “继续处理” picks up from the last finished step
    assert store.get_task(tid)['status'] == 'completed'


def test_result_states_workbook_and_locating_a_record(monkeypatch):
    mocked(monkeypatch)
    with TestClient(app) as client:
        tid = create(client, csv('期末.csv', FINAL), csv('平时.csv', USUAL), csv('补充.csv', EXTRA))
        wait(client, tid)
        draft = client.get(f'/api/tasks/{tid}/download/results.xlsx')
        assert draft.status_code == 200
        book = load_workbook(io.BytesIO(draft.content))
        assert book.sheetnames[:2] == ['说明（暂定）', '总评（暂定）'] and '问题清单' in book.sheetnames
        assert any(row[0].value == '待你确认' for row in book['问题清单'].iter_rows(min_row=2))
        task = decide(client, tid, client.get('/api/tasks/' + tid).json())
        book = load_workbook(io.BytesIO(client.get(f'/api/tasks/{tid}/download/results.xlsx').content))
        assert book.sheetnames == ['说明', '总评', '整合记录', '问题清单', '规则与依据', '材料目录']
        state = [row[1].value for row in book['说明'].iter_rows(min_row=2) if row[0].value == '结果状态'][0]
        assert state == '你已确认，结果完整'
        records = list(book['整合记录'].iter_rows(min_row=2, values_only=True))
        assert any('补充.csv 第2行' in (r[10] or '') for r in records)  # every record says where it came from
        source = task['result']['rows'][0]['sources'][0]
        found = client.get(f'/api/tasks/{tid}/locate', params={'file': source['file'], 'sheet': source['sheet'], 'row': source['row']}).json()
        assert found['kind'] == 'table' and any(r['row_number'] == source['row'] for r in found['rows'])
        assert client.get(f'/api/tasks/{tid}/locate', params={'file': '老师补录'}).status_code == 404


def test_a_confirmed_result_with_undecided_totals_says_so(monkeypatch):
    mocked(monkeypatch)
    with TestClient(app) as client:
        tid = create(client, csv('期末.csv', FINAL.replace('003,丙,期末,70', '003,丙,期末,缓考')), csv('平时.csv', USUAL))
        task = decide(client, tid, wait(client, tid))
        assert task['state']['code'] == 'confirmed_partial' and '1 人' in task['state']['text']
