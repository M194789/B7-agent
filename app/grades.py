from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import re
import unicodedata
from collections import Counter, defaultdict
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook

FIELDS = {
    'student_id': '学号', 'student_name': '姓名', 'course_id': '课程代码',
    'course_name': '课程名称', 'term': '学期', 'assessment': '考核项目',
    'occurrence': '次序/日期', 'attempt': '考试批次', 'score': '成绩',
    'max_score': '满分', 'status': '考试状态', 'remark': '备注',
    'student': '学号或姓名', 'question': '题号/问题', 'version': '提交批次/版本',
}
ALIASES = {
    'student_id': ['学号', '学生编号', '学籍号', 'student_id', 'studentid'],
    'student_name': ['姓名', '学生姓名', '学生', 'student_name', 'name'],
    'course_id': ['课程代码', '课程编号', 'course_id', 'coursecode'],
    'course_name': ['课程名称', '课程', '科目', 'course_name', 'course'],
    'term': ['学期', '学年学期', 'term', 'semester'],
    'assessment': ['考核项目', '考核类型', '项目', '记录类型', 'assessment', 'item'],
    'occurrence': ['次序', '第几次', '日期', '时间', '测验名称', '作业名称', 'occurrence', 'date'],
    'attempt': ['考试批次', '批次', '考试类型', 'attempt'],
    'score': ['成绩', '分数', '得分', '考试得分', 'score', 'grade'],
    'max_score': ['满分', '总分上限', 'max_score'], 'status': ['考试状态', '状态', '考试情况', 'status'],
    'remark': ['备注', '说明', 'remark'],
    'student': [], 'question': ['问题编号', '题号', '问题', 'question'],
    'version': ['提交批次', '提交次数', '提交版本', '版本', '第几次提交', 'version'],
}
# Ordered by precedence: a cheating mark outranks an absence mark on the same row.
STATUS_WORDS = {
    '作弊': ['作弊', '舞弊', '考试违纪'],
    '取消资格': ['取消考试资格', '取消资格'],
    '缓考': ['缓考', 'deferred'],
    '缺考': ['缺考', '旷考', '缺席考试', 'absent'],
    '免修': ['免修', '免考'],
    '未交': ['未交', '缺交', '未提交', '没交'],
    '待批改': ['待批改', '未批改', '待评阅', '待评分'],
}
MARKS = {'√', '✓', '✔', '到', '出勤', '是', '有', 'y', 'yes'}
ROSTER_ROLE = 'roster'
ITEM_WORDS = {'期中', '期末', '平时', '实验', '作业', '考勤', '小测', '测验', '期中考试', '期末考试', '平时成绩', '实验成绩'}


def text(value):
    return unicodedata.normalize('NFKC', str(value if value is not None else '')).strip()


def header_key(value):
    return re.sub(r'[\s_\-]+', '', text(value)).lower()


def schema_key(headers):
    return hashlib.sha256(json.dumps(headers, ensure_ascii=False).encode()).hexdigest()


def read_table(path, original_name=None):
    path = Path(path)
    name = original_name or path.name
    sheets = []
    if path.suffix.lower() == '.xlsx':
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in wb.worksheets:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append(row)
                    if len(rows) > 10001:
                        raise ValueError('单张表最多支持10000行，请拆分后上传。')
                if rows:
                    sheets.append((sheet.title, rows))
        finally:
            wb.close()
    elif path.suffix.lower() == '.xls':
        sheets = read_xls(path)
    else:
        decoded = decode_text(path.read_bytes())
        if decoded is None:
            raise ValueError('无法识别文件编码，请另存为 UTF-8 CSV。')
        try:
            dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=',\t;')
        except csv.Error:
            dialect = csv.excel_tab if path.suffix.lower() == '.tsv' else csv.excel
        rows = list(csv.reader(io.StringIO(decoded), dialect))
        if len(rows) > 10001:
            raise ValueError('单张表最多支持10000行。')
        sheets.append(('', rows))
    return build_tables(name, sheets)


class HtmlTables(HTMLParser):
    """Rows of every <table>; colspan cells keep their value in the first column like a merged Excel cell."""
    def __init__(self):
        super().__init__()
        self.tables, self.row, self.cell = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == 'table':
            self.tables.append([])
        elif tag == 'tr' and self.tables:
            self.row = []
        elif tag in ('td', 'th') and self.row is not None:
            span = str(dict(attrs).get('colspan') or '1')
            self.cell, self.span = [], int(span) if span.isdigit() and 0 < int(span) < 50 else 1

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        if tag in ('td', 'th') and self.cell is not None:
            self.row += [''.join(self.cell).strip()] + [''] * (self.span - 1)
            self.cell = None
        elif tag == 'tr' and self.row is not None:
            if any(self.row):
                self.tables[-1].append(self.row)
            self.row = None


def read_xls(path):
    """Old .xls. Exports from academic systems are often an HTML page renamed to .xls, so that is tried first."""
    raw = Path(path).read_bytes()
    if not raw.startswith(b'\xd0\xcf\x11\xe0'):
        decoded = decode_text(raw)
        if decoded and '<table' in decoded.lower():
            parser = HtmlTables()
            parser.feed(decoded)
            sheets = [(f'表{i + 1}' if len(parser.tables) > 1 else '', rows) for i, rows in enumerate(parser.tables) if len(rows) > 1]
            if sheets:
                return sheets
        raise ValueError('无法识别此 .xls 文件，请在 Excel 中另存为 .xlsx 后上传。')
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=raw)
    except xlrd.XLRDError:
        raise ValueError('无法打开此 .xls 文件，请在 Excel 中另存为 .xlsx 后上传。') from None
    sheets = []
    for sheet in book.sheets():
        if sheet.nrows > 10001:
            raise ValueError('单张表最多支持10000行，请拆分后上传。')
        rows = [[int(v) if isinstance(v, float) and v.is_integer() else v for v in sheet.row_values(i)] for i in range(sheet.nrows)]
        if rows:
            sheets.append((sheet.name, rows))
    return sheets


def decode_text(raw):
    for encoding in ('utf-8-sig', 'gb18030', 'utf-16'):
        try:
            return raw.decode(encoding)
        except UnicodeError:
            continue
    return None


SUMMARY_WORDS = ('合计', '总计', '平均', '平均分', '最高分', '最低分', '及格率', '优秀率', '统计', '制表', '审核', '任课教师签', '教师签')
KNOWN_HEADERS = {header_key(a) for aliases in ALIASES.values() for a in aliases} | {header_key(w) for w in ITEM_WORDS}


def is_number(value):
    try:
        float(text(value))
        return True
    except ValueError:
        return False


def header_score(row):
    cells = [text(c) for c in row if text(c)]
    known = sum(1 for c in cells if header_key(c) in KNOWN_HEADERS or split_item_header(c)[1] or split_item_header(c)[2] is not None)
    return known, len(cells), sum(1 for c in cells if is_number(c))


def value_like(cell):
    """A cell that reads as a value rather than a column name: a number, a date, a code such as Q3 or 00261003."""
    return is_number(cell) or bool(re.fullmatch(r'\d{1,4}\s*[/.\-月]\s*\d{1,2}日?|[A-Za-z]{0,3}\d+', cell))


def find_header(rows):
    """Exported score sheets often start with a title such as “2026秋 数据结构 成绩登记表”; the header is the first
    row among the top ten that names at least two known columns (or, failing that, the first wide row without numbers)."""
    top = rows[:10]
    width = max((len([c for c in r if text(c)]) for r in top), default=0)
    for index, row in enumerate(top):
        known, filled, numbers = header_score(row)
        if known >= 2 and numbers == 0:
            # “9/8, Q1, 张三” reads like column names too. When the rows above it are records (they hold values) under a row
            # just as wide that holds none (回答人 / 记录日期 / 问题), that row is the header and this one a record.
            if any(value_like(text(c)) for c in row if text(c)):
                for back in range(index - 1, -1, -1):
                    cells = [text(c) for c in top[back] if text(c)]
                    if not any(value_like(c) for c in cells):
                        if len(cells) == filled >= 2:
                            return back
                        break
            return index
    for index, row in enumerate(top):
        known, filled, numbers = header_score(row)
        if filled >= max(2, width * .6) and numbers == 0:
            return index
    return 0


def merged_headers(upper, lower):
    """Two header rows (“平时成绩” spanning “小测1 小测2”): fill the merged parent rightwards and prefer the sub-header."""
    parents, last = [], ''
    for cell in upper:
        last = str(cell).strip() if text(cell) else last
        parents.append(last)
    result = []
    for i in range(max(len(upper), len(lower))):
        parent = parents[i] if i < len(parents) else ''
        sub = text(lower[i]) if i < len(lower) else ''
        result.append(str(lower[i]).strip() if sub else parent)
    return result


def clean_headers(headers):
    seen, result = {}, []
    for index, h in enumerate(headers):
        h = str(h).strip() if text(h) else f'列{index + 1}'
        seen[h] = seen.get(h, 0) + 1
        result.append(h if seen[h] == 1 else f'{h}{seen[h]}')
    return result


def build_tables(name, sheets, uncertain=None, origin='file', locators=None):
    """Turn (sheet, rows) pairs into tables; `uncertain` maps a row number to columns an image reader was unsure about,
    `locators` maps a row number to where it sits in the original ('第4页第15行') when pages were joined into one table.

    For files the header row is located (title rows above it are kept as `preamble`), a two-row merged header is joined,
    and summary rows such as 平均分/合计 are set aside instead of being read as students."""
    tables = []
    for sheet, rows in sheets:
        start, notes, preamble = 0, [], []
        headers = [str(c) if c is not None else '' for c in rows[0]]
        if origin == 'file':
            start = find_header(rows)
            preamble = [' '.join(text(c) for c in r if text(c)) for r in rows[:start] if any(text(c) for c in r)]
            headers = [str(c) if c is not None else '' for c in rows[start]]
            nxt = rows[start + 1] if start + 1 < len(rows) else []
            if any(not text(h) for h in headers) and any(text(c) for c in nxt) and not any(is_number(c) for c in nxt) \
                    and header_score(nxt)[0] >= 1:
                headers = merged_headers(rows[start], nxt)
                notes.append(f'第 {start + 1}-{start + 2} 行是两级表头，已合并')
                start += 1
            if start and preamble:
                notes.append(f'跳过表头上方的标题行：{"；".join(preamble)[:120]}')
            raw = headers
            headers = clean_headers(headers)
            if headers != [str(h).strip() for h in raw]:
                notes.append('表头中有空白或重复列名，已自动编号，请在字段对应时核对')
        if len(headers) > 40:
            raise ValueError('单张表最多支持40列。')
        values, skipped = [], []
        for index, row in enumerate(rows[start + 1:], start + 2):
            if not any(text(c) for c in row):
                continue
            first = next((text(c) for c in row if text(c)), '')
            if any(first.startswith(w) for w in SUMMARY_WORDS):
                skipped.append(f'第{index}行“{first[:20]}”')
                continue
            item = {'row_number': index, 'values': {h: str(row[i]) if i < len(row) and row[i] is not None else '' for i, h in enumerate(headers)}}
            if uncertain and uncertain.get(index):
                item['uncertain'] = uncertain[index]
            if locators and locators.get(index):
                item['where'] = locators[index]
            values.append(item)
        if skipped:
            notes.append('统计/签名行不作为学生记录：' + '、'.join(skipped[:6]) + ('等' if len(skipped) > 6 else ''))
        table_id = hashlib.sha256((name + ':' + sheet).encode()).hexdigest()[:12]
        tables.append({'id': table_id, 'name': name, 'sheet': sheet, 'headers': headers, 'origin': origin,
                       'schema_key': schema_key(headers), 'rows': values, 'preamble': preamble, 'notes': notes})
    if not tables or not any(t['rows'] for t in tables):
        raise ValueError(f'{name} 没有可处理的数据行。')
    return [t for t in tables if t['rows']]


def read_roster(path, original_name=None):
    """Read a course roster. Only student ID and name columns are used."""
    roster = []
    for table in read_table(path, original_name):
        mapping = propose_mapping(table['headers'])
        if not (mapping.get('student_id') and mapping.get('student_name')):
            raise ValueError('课程名单需要包含“学号”和“姓名”两列。')
        for row in table['rows']:
            sid, name = text(row['values'][mapping['student_id']]), text(row['values'][mapping['student_name']])
            if sid:
                roster.append({'student_id': sid, 'student_name': name})
    if not roster:
        raise ValueError('课程名单中没有学生。')
    return roster


def propose_mapping(headers):
    result = {}
    for field, aliases in ALIASES.items():
        candidates = [h for h in headers if header_key(h) in {header_key(x) for x in aliases}]
        if len(candidates) == 1:
            result[field] = candidates[0]
    if 'score' not in result:
        # Platform exports often put each quiz/homework in its own column: "小测1", "作业2（10分）", "期中".
        items = [h for h in headers if h not in result.values() and (split_item_header(h)[1] or split_item_header(h)[2] is not None
                                                                     or text(h) in ITEM_WORDS)]
        if items:
            result['item_columns'] = items
    return result


def mapping_complete(mapping, request=None):
    request = request or {}
    identity = mapping.get('student_id') or mapping.get('student_name') or mapping.get('student')
    course = mapping.get('course_id') or mapping.get('course_name') or text(request.get('course'))
    # Columns named only by a date or week ('9/8', '第3周') need to be told which assessment they record.
    nameless = any(not split_item_header(c)[0] for c in mapping.get('item_columns') or [])
    scored = mapping.get('score') or mapping.get('item_columns') or mapping.get('total_column') or mapping.get('events')
    return bool(identity and course and scored and not (nameless and not text(mapping.get('item_name')))
                and not (mapping.get('events') and not text(mapping.get('item_name'))))


def validate_mapping(mapping, headers, request=None):
    result = {k: v for k, v in mapping.items() if k in FIELDS and isinstance(v, str) and v}
    items = [x for x in mapping.get('item_columns') or [] if isinstance(x, str) and x]
    total = mapping.get('total_column') if isinstance(mapping.get('total_column'), str) and mapping.get('total_column') else ''
    used = list(result.values()) + items + ([total] if total else [])
    if any(v not in headers for v in used):
        raise ValueError('字段映射引用了不存在的原始列。')
    if len(set(used)) != len(used):
        raise ValueError('同一原始列不能同时映射到多个标准字段。')
    if total:
        result['total_column'] = total
    if text(mapping.get('source_name')):
        result['source_name'] = text(mapping['source_name'])[:20]
    if mapping.get('events') is True:
        result['events'] = True
    if text(mapping.get('item_name')) and (items or result.get('events')):
        result['item_name'] = text(mapping['item_name'])[:40]
    if items:
        result['item_columns'] = items
    if not mapping_complete(result, request):
        raise ValueError('请至少指定学号或姓名、成绩列（或按列展开的成绩列），并指定课程列或在任务中填写课程名称；'
                         '成绩列只写了日期或周次时，还要填写这些列属于哪个考核项目（如“考勤”）。')
    return result


def split_item_header(header):
    """'小测2（10分）' -> ('小测', '2', 10.0). Used when each assessment is its own column.

    A column named only by a date or week ('9/8', '9月8日', '第3周') returns an empty assessment: the mapping's
    `item_name` (e.g. 考勤) supplies it. A weight note such as '平时(30%)' is dropped, it is not a full mark."""
    label, maximum = re.sub(r'[（(]\s*\d+(?:\.\d+)?\s*%\s*[)）]', '', text(header)).strip(), None
    found = re.search(r'[（(]\s*(?:满分)?\s*(\d+(?:\.\d+)?)\s*分?\s*[)）]', label)
    if found:
        maximum = float(found.group(1))
        label = (label[:found.start()] + label[found.end():]).strip()
    # Unit notes: '完成率(0-1)' is a fraction (×100), '(百分制)'/'(百分号)' are already out of 100, '(满分20)' above.
    unit = re.search(r'[（(]([^（）()]*)[)）]', label)
    if unit and maximum is None:
        note = unit.group(1)
        if re.search(r'(?<![\d.])0\s*[-~～到至]\s*1(?![\d.])', note):
            maximum = 1.0
        elif re.search(r'满分\s*(\d+(?:\.\d+)?)', note):
            maximum = float(re.search(r'满分\s*(\d+(?:\.\d+)?)', note).group(1))
        elif any(w in note for w in ('百分制', '百分号', '%', '百分比')):
            maximum = 100.0
        if maximum is not None:
            label = (label[:unit.start()] + label[unit.end():]).strip()
    date = re.fullmatch(r'(\d{1,2})\s*[/.月\-]\s*(\d{1,2})\s*日?', label)
    if date:
        return '', f'{int(date.group(1))}/{int(date.group(2))}', maximum
    week = re.fullmatch(r'第\s*(\d+)\s*[周次讲课]', label)
    if week:
        return '', week.group(1), maximum
    found = re.match(r'^(.*?)[\s_\-]*第?\s*(\d+)\s*次?$', label)
    if found and found.group(1):
        return found.group(1).strip(), found.group(2), maximum
    return label, '', maximum


def detect_status(*values):
    joined = ' '.join(text(v).lower() for v in values if v)
    for status, words in STATUS_WORDS.items():
        if any(w in joined for w in words):
            return status
    return ''


ID_LIKE = re.compile(r'[0-9A-Za-z\-]{5,}')


def issue_id(prefix, value):
    return prefix + ':' + hashlib.sha256(str(value).encode()).hexdigest()[:14]


def name_key(value):
    return re.sub(r'\s+', '', text(value))


def occurrence_key(value):
    value = text(value)
    return (0, float(value), '') if re.fullmatch(r'\d+(\.\d+)?', value) else (1, 0, value)


def locate(source):
    """Where a teacher finds the record: file, worksheet or page, and row — '考试成绩.pdf 第4页第15行', '平台.xlsx · 学习通 第4行'."""
    sheet = source.get('sheet') or ''
    if source.get('where'):
        return f"{source['file']} {source['where']}"
    return f"{source['file']}{' · ' + sheet if sheet else ''} 第{source['row']}行"


def shown(record):
    if record.get('status'):
        return record['status']
    if record.get('mark'):
        return '√（记录一次）'
    if record.get('grade'):
        return f"等级 {record['grade']}"
    return f"{record['score']:g} 分"


GRADE_WORDS = {'优秀', '良好', '中等', '及格', '不及格', '合格', '不合格', '优', '良', '中', '差'}


def grade_of(raw):
    """A letter or word grade (A, b+, 优秀); converted to a score later by the rule's own grade table, never guessed."""
    value = text(raw).upper()
    return value if re.fullmatch(r'[A-F][+-]?', value) or value in GRADE_WORDS else ''


EXAM_STATUSES = {'缺考', '缓考', '作弊', '取消资格'}
EXAM_ITEMS = ('期末', '期中', '考试', '卷面')
SOURCE_SEP = '·'


def item_assessment(column, mapping):
    """Assessment of one column in a wide table. Columns of a platform export carry the platform's name
    (“学习通·视频学习”) so that 视频学习 from 学习通 and 视频得分 from 慕课 never get mixed up; the platform's own
    total column becomes the platform itself (“学习通”)."""
    source = text(mapping.get('source_name'))
    if column == mapping.get('total_column'):
        return source or '平台总评'
    name, occurrence, _ = split_item_header(column)
    name = name or text(mapping.get('item_name')) or column
    # Numbered columns (小测1, Quiz 2) are repeated sittings of one assessment and keep its plain name, so the same quiz
    # exported by two platforms still merges; differently named columns (视频学习, 单元测验) are the platform's own parts.
    return f'{source}{SOURCE_SEP}{name}' if source and not occurrence else name


def row_note_status(row, mapping):
    """A status written in the row's status or remark column (not in a score cell)."""
    pick = lambda field: row['values'].get(mapping[field], '') if isinstance(mapping.get(field), str) else ''
    return detect_status(pick('status'), pick('remark'))


def note_applies(status, assessment, raw, exam_columns):
    """In a wide row a note such as “缺考” belongs to the exam column, not to the quizzes next to it.
    免修 covers the whole course; exam statuses cover the exam columns even when a score is written there (作弊 voids it),
    or the blank cells when the table has no exam column; anything else (未交) only fills blank cells."""
    if not status:
        return False
    if status == '免修':
        return True
    if status in EXAM_STATUSES and exam_columns:
        return any(w in assessment for w in EXAM_ITEMS)
    return not text(raw)


def has_exam_column(mapping):
    return any(any(w in item_assessment(c, mapping) for w in EXAM_ITEMS) for c in mapping.get('item_columns') or [])


def expand_rows(table, mapping):
    """One record per row in long tables; one record per filled cell of the item columns in wide platform exports
    (plus a blank cell that a status note in the same row refers to, e.g. an empty 期末 with “缺考” in 备注)."""
    items = list(mapping.get('item_columns') or [])
    if mapping.get('total_column') and mapping['total_column'] not in items:
        items.append(mapping['total_column'])
    exam_columns = has_exam_column(mapping)
    for row in table['rows']:
        base = f"{table['id']}:{row['row_number']}"
        if not items:
            yield base, base, row, None
            continue
        note = row_note_status(row, mapping)
        for index, column in enumerate(items):
            value = row['values'].get(column, '')
            if text(value) or note_applies(note, item_assessment(column, mapping), value, exam_columns):
                yield f'{base}:{index}', base, row, column


def version_number(record):
    found = re.search(r'\d+', record.get('version') or '')
    return int(found.group()) if found else 0


def latest_version(records):
    """Resubmissions of one assignment: the highest submission that has been graded counts; a newer one still waiting
    for grading (待批改) leaves the last graded one in place. Different grades within that same submission stay a
    conflict for the teacher — the older version is never used to hide it. Returns (records kept, records superseded)."""
    if len({version_number(r) for r in records}) < 2:
        return records, []
    graded = [r for r in records if r['status'] != '待批改']
    top = max(version_number(r) for r in graded or records)
    return [r for r in records if version_number(r) == top], [r for r in records if version_number(r) != top]


def merge(tables, mappings, request=None, decisions=None, roster=None):
    request, decisions = request or {}, decisions or {}
    roster_ids = {r['student_id']: r['student_name'] for r in roster or []}
    roster_names = defaultdict(list)
    for r in roster or []:
        roster_names[name_key(r['student_name'])].append(r['student_id'])
    groups = defaultdict(list)
    issues, excluded, total, name_matched = [], [], 0, 0
    for table in tables:
        mapping = mappings[table['id']]
        exam_columns = has_exam_column(mapping)
        def get(row, field):
            column = mapping.get(field)
            return row['values'].get(column, '') if isinstance(column, str) else ''
        for rid, base, row, column in expand_rows(table, mapping):
            total += 1
            if column:
                _, occurrence, column_max = split_item_header(column)
                assessment = item_assessment(column, mapping)
                raw = row['values'][column]
            elif mapping.get('events'):
                # A log where every row is one event (an answer, an attendance): it counts once; the date and question
                # number tell repeated exports of the same answer apart from different answers on the same day.
                assessment = text(mapping.get('item_name')) or '记录'
                occurrence = ' '.join(x for x in (text(get(row, 'occurrence')), text(get(row, 'question'))) if x) or f'第{row["row_number"]}行'
                column_max, raw = None, '√'
            else:
                assessment = text(get(row, 'assessment')) or request.get('assessment', '期末')
                occurrence = ' '.join(x for x in (text(get(row, 'occurrence')), text(get(row, 'question'))) if x)
                column_max, raw = None, get(row, 'score')
            sid, cid, cname = text(get(row, 'student_id')), text(get(row, 'course_id')), text(get(row, 'course_name'))
            name_written = text(get(row, 'student_name'))
            either = text(get(row, 'student'))
            if either and not (sid or name_written):
                # One column holds a student number or a name, as on a hand-written answer sheet; the roster tells which.
                if either in roster_ids or (re.fullmatch(r'[0-9A-Za-z\-]{4,}', either) and re.search(r'\d', either)):
                    sid = either
                else:
                    name_written = either
            assumed = not (cid or cname), not text(get(row, 'term'))  # filled in from the task, not read from the record
            if assumed[0]:
                cname = text(request.get('course'))
            record = {'id': rid, 'student_id': sid, 'student_name': name_written, 'assumed': list(assumed),
                      'course_id': cid, 'course_name': cname, 'term': text(get(row, 'term')) or request.get('term', '2026秋'),
                      'assessment': assessment, 'occurrence': occurrence,
                      'attempt': text(get(row, 'attempt')) or request.get('attempt', '初考'),
                      'version': text(get(row, 'version')),
                      'remark': get(row, 'remark'), 'raw_score': raw, 'status': '', 'mark': False, 'matched_by': '',
                      'source': {'file': table['name'], 'sheet': table['sheet'], 'row': row['row_number'], 'column': column or '',
                                 'where': row.get('where', ''), 'origin': table.get('origin', 'file'), 'raw': row['values']}}
            record['assessment'] = {'final': '期末', 'coursework': '平时', '期末考试': '期末', '平时成绩': '平时', '期中考试': '期中'}.get(record['assessment'].lower(), record['assessment'])
            record['attempt'] = {'first': '初考', 'resit': '补考', '正考': '初考'}.get(record['attempt'].lower(), record['attempt'])
            # Identity is decided once per source row; in wide tables each cell's value is decided separately.
            who = decisions.get('row:' + base, {})
            what = decisions.get('row:' + rid, {}) if rid != base else who
            if 'exclude' in (who.get('action'), what.get('action')):
                excluded.append({**record, 'reason': who.get('note') or what.get('note') or '人工排除'})
                continue
            if who.get('action') == 'correct' and text(who.get('student_id')):
                record['student_id'] = sid = text(who['student_id'])
            action = who.get('action')
            corrected = what.get('action') == 'correct' and bool(text(what.get('score')))
            if corrected:
                raw = what['score']
            # A teacher's corrected value replaces whatever the status column or remark said.
            notes = (get(row, 'status'), record['remark'])
            if column and not detect_status(raw) and not note_applies(row_note_status(row, mapping), assessment, raw, exam_columns):
                notes = ()  # the note in this wide row is about another column
            record['status'] = detect_status(raw) if corrected else detect_status(raw, *notes)
            # 缺考 and 缓考 usually lead to different outcomes (zero vs pending), so mixed signals go to the teacher.
            mixed = not corrected and {'缺考', '缓考'} <= {detect_status(v) for v in (raw, *notes)}
            score = maximum = None
            invalid = '成绩或满分无效（缺失不等于零分）'
            if not record['status']:
                if text(raw).lower() in MARKS:
                    record['mark'] = True
                elif grade_of(raw):
                    record['grade'] = grade_of(raw)
                else:
                    percent_sign = re.fullmatch(r'(-?\d+(?:\.\d+)?)\s*%', text(raw))  # '85%' is 85 out of 100
                    if not text(raw):
                        invalid = '成绩为空（不会当作 0 分）'
                    elif not (re.fullmatch(r'-?\d+(\.\d+)?', text(raw)) or percent_sign):
                        invalid = f'成绩“{text(raw)[:20]}”不是数字或等级'
                    try:
                        value = float(percent_sign.group(1) if percent_sign else text(raw))
                        maximum = 100.0 if percent_sign else float(text(get(row, 'max_score'))) if text(get(row, 'max_score')) else column_max
                        if corrected and maximum is not None and maximum <= 0:
                            maximum = None  # the teacher typed the real score; a broken full mark (0) must not keep blocking it
                        limit = maximum if maximum is not None else 100
                        if maximum is not None and maximum <= 0:
                            invalid = f'满分写的是 {maximum:g}，无法换算'
                        if not (math.isfinite(limit) and math.isfinite(value) and limit > 0):
                            raise ValueError('无效数值')
                        if not 0 <= value <= limit:
                            invalid = f'分数 {value:g} 超出范围（满分 {limit:g}{"，未注明满分时按100" if maximum is None else ""}）'
                            raise ValueError(invalid)
                        score = round(value / limit * 100, 4)
                        record['value'] = value
                    except (ValueError, TypeError, OverflowError):
                        score = None
            record['max_score'] = maximum
            problems, values, candidates = [], [], []
            if action != 'accept' or not sid:
                if not sid:
                    name = name_key(record['student_name'])
                    if not name:
                        problems.append('缺少学号和姓名')
                    elif not roster:
                        problems.append('只有姓名、没有学号，且未提供课程名单，无法确认是哪位学生')
                    elif len(roster_names.get(name, [])) == 1:
                        record['student_id'] = sid = roster_names[name][0]
                        record['matched_by'] = 'roster_name'
                        name_matched += 1
                    elif roster_names.get(name):
                        candidates = roster_names[name]
                        problems.append(f'课程名单中有 {len(candidates)} 名同名学生，原记录未写学号，请指定是哪一位')
                    else:
                        problems.append('课程名单中没有这个姓名')
                elif roster:
                    if sid not in roster_ids:
                        problems.append('学号不在课程名单中（可能录错或不属于本课程）')
                        candidates = roster_names.get(name_key(record['student_name']), [])
                    elif record['student_name'] and name_key(record['student_name']) != name_key(roster_ids[sid]) and action != 'correct':
                        # When the teacher picked the student, “张伟（1班）” versus “张伟” is not asked again.
                        problems.append(f'学号与姓名不一致：名单中该学号是“{roster_ids[sid]}”')
                        # Either the number or the name is wrong: both students are offered, and both wait.
                        candidates = [sid] + [c for c in roster_names.get(name_key(record['student_name']), []) if c != sid]
                # A blank cell the reader marked unsure has nothing to check: a missing 学号 is resolved by the name,
                # a missing score is reported as blank on its own.
                filled = lambda c: bool(text(row['values'].get(c)))
                unsure = [c for c in row.get('uncertain', []) if c in (mapping.get('student_id'), mapping.get('student_name'), mapping.get('student')) and filled(c)]
                if unsure and action != 'correct':
                    problems.append('识别不确定：' + '、'.join(unsure) + '，请对照原件核对')
            if not (record['course_id'] or record['course_name']):
                problems.append('缺少课程代码和名称')
            unsure = [c for c in row.get('uncertain', []) if c == (column or mapping.get('score')) and text(row['values'].get(c))]
            if unsure and what.get('action') not in ('correct', 'accept'):
                values.append('识别不确定：' + '、'.join(unsure) + '，请对照原件核对')
            if score is None and not record['status'] and not record['mark'] and not record.get('grade'):
                values.append(invalid)
            if mixed and what.get('action') != 'accept':
                values.append(f'同一条记录同时出现“缺考”和“缓考”（成绩：{text(raw)}；备注：{text(record["remark"])}），请在“修正成绩”中填写实际状态')
            where = f"{locate(record['source'])} · {record['student_name'] or '姓名未提供'}"
            found = [(problems, base, 'identity' if column else 'row'), (values, rid, 'value' if column else 'row')]
            if not column:
                found = [(problems + values, rid, 'row')]
            for listed, key_id, scope in found:
                if listed and not any(i['id'] == 'row:' + key_id for i in issues):
                    issues.append({'id': 'row:' + key_id, 'type': 'row', 'scope': scope, 'title': ' / '.join(listed),
                                   'description': where + (f' · {column}（原值 {raw}）' if scope == 'value' else ''),
                                   'record': record, 'candidates': [{'student_id': c, 'student_name': roster_ids.get(c, '')} for c in candidates], 'options': []})
            if problems or values:
                continue
            record['score'] = record['normalized_score'] = score
            key = (record['student_id'], record['course_id'] or 'name:' + record['course_name'], record['term'],
                   record['assessment'], record['occurrence'], record['attempt'])
            groups[key].append(record)
    output, conflicts, duplicate_rows, duplicates = [], 0, 0, []
    compare = lambda r: 'status:' + r['status'] if r['status'] else 'mark' if r['mark'] else 'grade:' + r['grade'] if r.get('grade') else r['score']
    for key, records in groups.items():
        iid = issue_id('conflict', key)
        records, superseded = latest_version(records)
        counts = Counter(compare(r) for r in records)
        duplicate_rows += sum(n - 1 for n in counts.values())
        if any(n > 1 for n in counts.values()):
            duplicates.append({'student_id': key[0], 'course': records[0]['course_name'] or key[1], 'key': list(key),
                               'duplicate_count': sum(n - 1 for n in counts.values()), 'sources': [r['source'] for r in records]})
        # Identical IDs with genuinely different names require review even if scores agree.
        names = {name_key(r['student_name']) for r in records if r['student_name']}
        identity_conflict = len(names) > 1
        choice = decisions.get(iid, {})
        if len(counts) > 1 or identity_conflict:
            conflicts += 1
            if choice.get('action') == 'exclude':
                excluded.extend({**r, 'reason': '人工排除冲突组'} for r in records)
                continue
            if choice.get('action') == 'select':
                selected = next((r for r in records if r['id'] == choice.get('record_id')), None)
                if selected is None:
                    raise ValueError('选中的记录不属于当前冲突组。')
            else:
                label = f"{key[3]}{' 第' + key[4] + '次' if re.fullmatch(r'\d+', key[4]) else ' ' + key[4] if key[4] else ''} / {key[5]}"
                issues.append({'id': iid, 'type': 'conflict', 'title': '同一学号的姓名存在差异' if identity_conflict else '同一考核记录在不同来源中不一致',
                               'description': f"{key[0]} · {records[0]['course_name'] or key[1]} · {label}",
                               'options': [{'value': r['id'], 'label': f"{shown(r)} · {r['student_name']} · {locate(r['source'])}"} for r in records],
                               'records': records})
                continue
        else:
            selected = records[0]
        output.append({**selected, 'sources': [r['source'] for r in records],
                       'reviewed': bool(choice), 'review_note': choice.get('note', ''), 'source_count': len(records),
                       'superseded': [{'version': r['version'], 'shown': shown(r), 'source': r['source']} for r in superseded]})
    output, suspected = possible_duplicates(output, decisions)
    issues += suspected
    output.sort(key=lambda r: (r['student_id'], r['course_id'] or r['course_name'], r['assessment'], occurrence_key(r['occurrence']), r['attempt']))
    present = {r['student_id'] for r in output} | {i['record']['student_id'] for i in issues if i['type'] == 'row'}
    missing = [{'student_id': sid, 'student_name': name} for sid, name in roster_ids.items() if sid not in present]
    rules = ['按学号、课程代码（缺失时用课程名）、学期、考核项目、次序/日期、考试批次组成业务键：同一次小测从两个平台导出会合并，第1次与第2次小测分别保留。',
             '同一业务键、同一成绩归并并保留全部来源；不同成绩或状态等待人工选择。',
             '缺考、缓考、作弊、免修等作为考试状态单独记录，不当作分数，也不当作错误数据。',
             '学号保留前导零与连接符；平时与期末、初考与补考分别保留。']
    if roster:
        rules.append('提供了课程名单：只写姓名的记录按名单唯一匹配，同名或查无此人交由老师确认；学号不在名单、学号与姓名不符均提示。')
    return {'kind': 'grades', 'rows': output, 'issues': issues, 'duplicates': duplicates, 'excluded': excluded,
            'missing_students': missing,
            'summary': {'input_rows': total, 'output_rows': len(output), 'duplicate_rows': duplicate_rows,
                        'conflict_groups': conflicts, 'pending': len(issues), 'excluded_rows': len(excluded),
                        'students': len({r['student_id'] for r in output}), 'status_rows': sum(1 for r in output if r['status']),
                        'name_matched': name_matched, 'missing_students': len(missing), 'roster_size': len(roster_ids)},
            'rules': rules}


def possible_duplicates(output, decisions):
    """The same answer logged twice in two sources, once as “9/8 Q3” and once only as “9/8”: the keys differ, so the
    records were not merged, yet they may be one event. Records of one student and item from different files, where one
    occurrence says everything the other says and more, are put to the user; nothing is merged or dropped on a guess.
    Different questions, homework or exams (“9/8 Q3” and “9/8 Q4”) never qualify."""
    groups = defaultdict(list)
    for r in output:
        if r['student_id']:
            groups[(r['student_id'], r['course_id'] or r['course_name'], r['term'], r['assessment'], r['attempt'])].append(r)
    dropped, issues = set(), []
    files = lambda r: {(x['file'], x['sheet']) for x in r['sources']}
    tokens = lambda r: set(r['occurrence'].split())
    for records in groups.values():
        for short in records:
            # Records in other files that say all this one says and more (“9/8” -> “9/8 Q3”, “9/8 Q4”).
            fuller = [r for r in records if r is not short and r['id'] not in dropped and tokens(short) < tokens(r) and not files(r) & files(short)]
            if not fuller or short['id'] in dropped:
                continue
            iid = issue_id('dup', short['id'])
            choice = decisions.get(iid, {})
            pairs = {f'{keep["id"]}|{drop["id"]}': (keep, drop) for r in fuller for keep, drop in ((r, short), (short, r))}
            if choice.get('action') == 'select' and choice.get('record_id') in pairs:
                keep, drop = pairs[choice['record_id']]
                keep['sources'] = keep['sources'] + drop['sources']
                keep['source_count'] = len(keep['sources'])
                dropped.add(drop['id'])
            elif choice.get('action') != 'keep':
                options = []
                for r in fuller:
                    if shown(r) == shown(short):
                        options.append({'value': f'{r["id"]}|{short["id"]}', 'label': f'是同一次：就是“{r["occurrence"]}”那条（{locate(r["sources"][0])}），只计一次'})
                    else:
                        options += [{'value': f'{keep["id"]}|{drop["id"]}', 'label': f'是“{r["occurrence"]}”那一次，成绩按 {shown(keep)}（{locate(keep["sources"][0])}）'}
                                    for keep, drop in ((r, short), (short, r))]
                issues.append({'id': iid, 'type': 'duplicate',
                               'title': '这条记录可能和另一份材料中的记录是同一次' + ('' if all(shown(r) == shown(short) for r in fuller) else '（成绩不同）'),
                               'description': f"{short['student_id']} {short['student_name']} · {short['assessment']}：{locate(short['sources'][0])} 写的是“{short['occurrence'] or '未写次序'}”，"
                                              f"另有 {'、'.join('“' + r['occurrence'] + '”' for r in fuller)}（{locate(fuller[0]['sources'][0])}）",
                               'options': options, 'records': [short] + fuller})
    return [r for r in output if r['id'] not in dropped], issues


def safe_cell(value):
    value = str(value)
    return "'" + value if value.startswith(('=', '+', '-', '@', '\t', '\r')) else value


def csv_bytes(result):
    stream = io.StringIO(newline='')
    writer = csv.writer(stream)
    writer.writerow(['学号', '姓名', '课程代码', '课程名称', '学期', '考核项目', '次序/日期', '考试批次', '百分制成绩', '考试状态', '来源数', '来源位置', '人工复核'])
    for r in result.get('rows', []):
        score = '' if r['score'] is None else r['score']
        values = [r['student_id'], r['student_name'], r['course_id'], r['course_name'], r['term'], r['assessment'], r['occurrence'], r['attempt'],
                  score, r['status'] or ('记录一次' if r['mark'] else ''), r['source_count'],
                  '; '.join(f"{s['file']}#{s['row']}{'/' + s['column'] if s.get('column') else ''}" for s in r['sources']), '是' if r['reviewed'] else '否']
        writer.writerow([safe_cell(x) for x in values])
    return stream.getvalue().encode('utf-8-sig')


