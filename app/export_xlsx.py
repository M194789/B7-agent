"""The result as one workbook a teacher can open: totals or scores, the issue list, the rules and how they were
applied, the materials, and every record with where it came from. Also used to show a record in its original file.

A workbook taken while the task still waits for confirmation says so in its title and on every sheet name it
matters to; it is never presented as a final result.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import appraisal, grades, scoring, store, versions
from .documents import IMAGE_SUFFIXES

HEAD = PatternFill('solid', fgColor='E8F1EF')
WARN = PatternFill('solid', fgColor='FFF4DD')
STATE_WORDS = {'ok': '已计算', 'pending': '待定', 'invalid': '无效', 'exempt': '免修'}


def cell(value):
    if value is None:
        return ''
    if isinstance(value, (int, float)):
        return value
    return grades.safe_cell(value)


def sheet(book, title, headers, rows, widths=None, warn=None):
    ws = book.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HEAD
    for row in rows:
        ws.append([cell(v) for v in row])
        if warn and warn(row):
            for c in ws[ws.max_row]:
                c.fill = WARN
    for i, width in enumerate(widths or [], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    return ws


def where(source):
    return grades.locate(source) + (f' · {source["column"]}' if source.get('column') else '')


def build(task):
    """The workbook bytes, from the confirmed result or, while waiting, from what the confirmation page shows."""
    book = Workbook()
    book.remove(book.active)
    (grades_book if task['kind'] == 'grades' else appraisal_book)(book, task)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def write(task):
    output = store.DATA / 'tasks' / task['id'] / 'output'
    output.mkdir(parents=True, exist_ok=True)
    (output / 'results.xlsx').write_bytes(build(task))


def about(book, task, lines):
    state = versions.result_state(task)
    final = task['status'] == 'completed'
    rows = [('任务', task['name']), ('结果状态', state['text'] + ('' if final else '（暂定，不是最终结果）')),
            ('版本', f'第 {versions.number(task)} 版'), ('导出时间', store.now()),
            ('模型', '有实际模型调用（语义理解、识别）；所有分数由程序计算' if task['llm_used'] else '没有调用模型')] + lines
    ws = sheet(book, '说明' if final else '说明（暂定）', ['项目', '内容'], rows, [16, 100])
    if not final:
        ws['B2'].fill = WARN


def change_rows(changes):
    if not changes:
        return []
    rows = []
    for key, label in (('files_added', '新增材料'), ('files_replaced', '替换材料'), ('files_removed', '删除材料')):
        if changes.get(key):
            rows.append((label, '、'.join(changes[key])))
    if changes.get('rules_changed'):
        rows.append(('规则', '改了规则原文，按新规则重新理解'))
    if changes.get('was_confirmed'):
        rows.append(('上一版', f'第 {changes.get("previous")} 版已确认的结果已被这一版取代'))
    for x in changes.get('lapsed') or []:
        rows.append(('处理已失效', f'{x.get("title", "")}（原来：{x["label"]}）'))
    for x in changes.get('revoked') or []:
        rows.append(('撤回的处理', f'{x["title"]}（原来：{x["label"]}）'))
    for t in changes.get('totals_moved') or []:
        before = '无' if not t['before'] else t['before'][0] if t['before'][0] is not None else STATE_WORDS.get(t['before'][1], '待定')
        after = t['after'][0] if t['after'][0] is not None else STATE_WORDS.get(t['after'][1], '待定')
        rows.append(('总评变化', f'{t["student_id"]} {t["student_name"]}：{before} → {after}'))
    return rows


def grades_book(book, task):
    result = task['result']
    pending = result.get('pending') if task['status'] == 'waiting_review' else None
    view = pending or result
    rule = view.get('rule')
    totals = (pending.get('preview') if pending else result.get('totals')) or []
    counts = {k: sum(1 for t in totals if t['outcome'] == k) for k in STATE_WORDS}
    lines = [('材料', f'{len(task["files"])} 份'), ('整合记录', f'{len(result.get("rows") or [])} 条')]
    if rule:
        lines.append(('总评', f'已算出 {counts["ok"]} 人；待定 {counts["pending"]} 人；无效 {counts["invalid"]} 人；免修 {counts["exempt"]} 人'))
    lines += change_rows(view.get('changes'))
    about(book, task, lines)
    if rule:
        comps = rule['components']
        rows = []
        for t in totals:
            hows = '；'.join(f'{p["name"]}：{p["how"]}' for p in t['components'] if p.get('how'))
            rows.append([t['student_id'], t['student_name'], t.get('course', ''), t.get('term', ''),
                         *[p['percent'] for p in t['components']], t['total'], STATE_WORDS[t['outcome']], '；'.join(t['notes']), hows])
        sheet(book, '总评' if not pending else '总评（暂定）', ['学号', '姓名', '课程', '学期', *[f'{c["name"]}（{c["weight"]:g}%）' for c in comps],
                                                             '总评', '状态', '说明', '计算过程'],
              rows, [14, 10, 14, 12, *[12] * len(comps), 9, 8, 40, 60], warn=lambda r: r[4 + len(comps) + 1] != '已计算')
    sheet(book, '整合记录', ['学号', '姓名', '课程代码', '课程', '学期', '考核项目', '次序/日期', '批次', '成绩/状态', '原始值', '来源', '人工处理'],
          [[r['student_id'], r['student_name'], r['course_id'], r['course_name'], r['term'], r['assessment'], r['occurrence'], r['attempt'],
            grades.shown(r), grades.text(r.get('raw_score')), '\n'.join(where(s) for s in r['sources']), '是' if r.get('reviewed') else '']
           for r in result.get('rows') or []], [14, 10, 12, 14, 12, 16, 14, 8, 12, 10, 50, 8])
    sheet(book, '问题清单', ['状态', '问题', '位置与说明', '处理'], issue_rows(view, totals, pending), [14, 44, 60, 40],
          warn=lambda r: r[0] in ('待你确认', '已失效', '请核对'))
    left = [[r['student_id'], r['student_name'], r['assessment'], r['occurrence'], grades.shown(r) if r.get('score') is not None or r.get('status') else grades.text(r.get('raw_score')),
             r.get('reason', ''), where(r['source'])] for r in result.get('excluded') or []]
    left += [['', '', '', '', '', '不属于本次统计范围', o['text']] for o in view.get('outside') or []]
    if left:
        sheet(book, '未计入的记录', ['学号', '姓名', '考核项目', '次序', '成绩', '原因', '来源'], left, [14, 10, 16, 12, 10, 24, 50])
    basis = []
    if view.get('source_text') or (rule or {}).get('source_text'):
        basis.append(('规则原文', view.get('source_text') or rule.get('source_text')))
    basis += [('采用的规则', x) for x in (view.get('rule_lines') or (scoring.describe(rule) if rule else []))]
    basis += [('按你的话修改', f'“{r["text"]}”') for r in view.get('revisions') or []]
    plan = view.get('plan') or {}
    basis += [('开始前的说明', f'{q["question"]} → {a}') for q, a in zip(plan.get('questions') or [], plan.get('answers') or []) if a]
    if plan.get('baseline_note'):
        basis.append(('完整性依据', plan['baseline_note']))
    basis += [('处理口径', x) for x in result.get('rules') or []]
    basis += [('材料的理解', f'{t["name"]}（{t["rows"]} 行）：{t["text"]}') for t in view.get('tables') or []]
    sheet(book, '规则与依据', ['类别', '内容'], basis, [16, 110])
    statuses = {'imported': '已导入', 'roster': '名单', 'rules': '评分规则', 'duplicate': '重复，已忽略', 'empty': '没有成绩', 'skipped': '未处理', 'failed': '没读出来'}
    sheet(book, '材料目录', ['材料', '状态', '说明'], [[m['name'], statuses.get(m['status'], m['status']), m['text']] for m in view.get('materials') or []],
          [50, 14, 60], warn=lambda r: r[1] in ('没读出来', '未处理', '没有成绩'))


def issue_rows(view, totals, pending):
    rows = []
    for i in (pending or {}).get('issues') or []:
        s = i.get('suggestion')
        rows.append(['已失效' if i.get('was') else '待你确认', i['title'], i['description'],
                     (f'原来的处理“{i["was"]}”因材料变化失效；' if i.get('was') else '') + (f'建议：{s["label"]}' if s else '需要你选择')])
    for x in view.get('decided') or []:
        rows.append(['已处理', x['title'], x['description'], x['label']])
    for f in view.get('failed_files') or []:
        rows.append(['请核对', f'{f["name"]} 没有读出来', f['error'], '其中的成绩没有计入'])
    for z in view.get('zeros') or []:
        for item in z['items']:
            rows.append(['请核对' if not item.get('pending') else '待定', f'{z["student_id"]} {z["student_name"]}：{item["text"]}', '没有记录或按规则记 0 分', ''])
    for m in view.get('mismatches') or []:
        rows.append(['请核对', f'{m["student_id"]} {m["student_name"]}：平台总评与分项重算不一致', '；'.join(m['items']), '按重算值计'])
    for x in view.get('remark_commands') or []:
        rows.append(['已忽略', '备注里像算分指令的文字', x, '只当作数据，没有照做'])
    for m in view.get('missing_students') or []:
        rows.append(['请核对', f'{m["student_id"]} {m["student_name"]} 没有任何记录', '名单中有此人', ''])
    for t in totals:
        if t['outcome'] != 'ok':
            rows.append(['待定' if t['outcome'] == 'pending' else STATE_WORDS[t['outcome']], f'{t["student_id"]} {t["student_name"]} 的总评{STATE_WORDS[t["outcome"]]}',
                         '；'.join(t['notes']), '不是正式成绩'])
    for w in view.get('warnings') or []:
        rows.append(['请核对', w, '', ''])
    return rows


def appraisal_book(book, task):
    result = task['result']
    pending = result.get('pending') if task['status'] == 'waiting_review' else None
    view = pending or result
    indicators = view['scheme']['indicators']
    people = (pending.get('preview') if pending else result.get('results')) or []
    about(book, task, [('被考核人', f'{len(people)} 人'), ('材料', f'{len(view.get("materials") or [])} 份')] + change_rows(view.get('changes')))
    sheet(book, '汇总' if not pending else '汇总（暂定）', ['被考核人', *[f'{i["name"]}（{i["max"]:g}）' for i in indicators], '总分', '状态'],
          [[r['person'], *[i['score'] for i in r['items']], r['total'] if r['total'] is not None else r.get('partial'),
            '待评定' if r.get('pending') else '已完成'] for r in people], [12, *[12] * len(indicators), 9, 9], warn=lambda r: r[-1] != '已完成')
    names = {m['id']: m['file'] for m in view.get('materials') or []}
    sheet(book, '评分依据', ['被考核人', '指标', '得分', '满分', '计分方式', '依据', '材料'],
          [[r['person'], f'{i["group"]} / {i["name"]}', i['score'], i['max'], i['method'], i['how'], '、'.join(names.get(m, m) for m in i['materials'])]
           for r in people for i in r['items']], [12, 30, 8, 8, 12, 60, 40])
    if pending:
        sheet(book, '待评分', ['被考核人', '指标', '满分', '建议', '材料'],
              [[j['person'], j['name'], j['max'], (j.get('suggestion') or {}).get('level', ''), '、'.join(names.get(m, m) for m in j['materials'])]
               for j in pending.get('judge') or [] if not j.get('decided')], [12, 30, 8, 16, 50])
    by_id = {i['id']: i['name'] for i in indicators}
    states = appraisal.STATUSES
    sheet(book, '材料目录', ['材料', '状态', '原因', '被考核人', '指标', '类型', '级别', '日期', '角色', '内容'],
          [[m['file'], states.get(m.get('status'), ''), '；'.join(m.get('reasons') or []), m.get('person', ''), by_id.get(m.get('indicator'), ''),
            m.get('doc_type', ''), m.get('level', ''), m.get('date', ''), m.get('role', ''), m.get('summary', '')] for m in view.get('materials') or []],
          [36, 12, 50, 10, 20, 12, 8, 12, 12, 50], warn=lambda r: r[1] in ('需要补充证明', '不符合条件'))
    import_lines = [('指标', x) for x in view.get('scheme_lines') or []]
    if view.get('source_text') or view['scheme'].get('source_text'):
        import_lines.insert(0, ('方案原文', view.get('source_text') or view['scheme']['source_text']))
    sheet(book, '考核方案', ['类别', '内容'], import_lines, [12, 110])


PAGE = re.compile(r'第(\d+)(?:-\d+)?页')


def locate(file, index, sheet_name, row, where_text):
    """The original around one record: rows of a table, or the page or photo it was read from (positions inside a
    photo are not known, so the photo is shown whole with the recognised row next to it)."""
    path, name = file['path'], file['name']
    suffix = Path(path).suffix.lower()
    url = f'sources/{index}?inline=1'
    if suffix in IMAGE_SUFFIXES:
        return {'kind': 'image', 'url': url, 'file': name, 'where': where_text}
    if suffix == '.pdf':
        page = PAGE.search(where_text or '') or PAGE.search(sheet_name or '')
        return {'kind': 'pdf', 'url': url + (f'#page={page.group(1)}' if page else ''), 'file': name, 'where': where_text,
                'page': int(page.group(1)) if page else None}
    try:
        if suffix == '.docx':
            from .workflow import read_word_tables
            tables = read_word_tables(path, name)
        else:
            tables = grades.read_table(path, name)
    except ValueError as error:
        return {'kind': 'file', 'url': url, 'file': name, 'error': str(error)[:200]}
    table = next((t for t in tables if (t.get('sheet') or '') == (sheet_name or '')), tables[0] if tables else None)
    if not table:
        return {'kind': 'file', 'url': url, 'file': name}
    near = [r for r in table['rows'] if abs(r['row_number'] - row) <= 3]
    return {'kind': 'table', 'file': name, 'sheet': table.get('sheet') or '', 'headers': table['headers'], 'row': row,
            'rows': [{'row_number': r['row_number'], 'values': [r['values'].get(h, '') for h in table['headers']]} for r in near],
            'url': f'sources/{index}'}
